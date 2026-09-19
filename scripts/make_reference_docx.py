"""Generate ``reports/reference.docx`` used by pandoc for the report.

Applies GOST 7.32-2017 / ITMO formatting: Times New Roman 14 pt, 1.5 line
spacing, first-line indent 1.25 cm, justified body, margins 3/1/2/2 cm and page
numbers. Run: ``python scripts/make_reference_docx.py``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path("reports/reference.docx")


def _set_page(document: Document) -> None:
    section = document.sections[0]
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def _set_normal(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(0)


def _add_page_numbers(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = run._r.makeelement(qn("w:fldChar"), {})
    begin.set(qn("w:fldCharType"), "begin")
    instr = run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    instr.text = "PAGE"
    end = run._r.makeelement(qn("w:fldChar"), {})
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, end):
        run._r.append(element)


def main() -> None:
    """Build and save the reference document."""
    document = Document()
    _set_page(document)
    _set_normal(document)
    _add_page_numbers(document)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
